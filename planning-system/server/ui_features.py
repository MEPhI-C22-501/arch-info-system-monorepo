"""Project cards, resource choices and user XLSX exports (separate from Reports ETL)."""
import copy
import math
import os
from datetime import date
from io import BytesIO

from fastapi import Request
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.cell.cell import Cell, ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill

from . import storage as db
from .auth import actor
from .domain import fail, validate_draft

PROJECT_STATUSES = {'DRAFT', 'ACTIVE', 'ON_HOLD', 'COMPLETED', 'ARCHIVED'}
ALLOCATION_STATUSES = {'PROPOSED', 'CONFIRMED', 'CANCELLED'}


def filters(req, assignment=False):
    permitted = {'q', 'projectId', 'portfolioId', 'status', 'dateFrom', 'dateTo'}
    if assignment:permitted.add('employeeId')
    if set(req.query_params)-permitted:fail(422, 'Неизвестный фильтр')
    if any(len(req.query_params.getlist(k)) != 1 for k in req.query_params):fail(422, 'Повтор фильтра')
    values = dict(req.query_params)
    for key in ['dateFrom', 'dateTo']:
        if values.get(key):
            try:
                parsed = date.fromisoformat(values[key])
                if parsed.isoformat() != values[key]:raise ValueError()
            except ValueError:fail(422, 'Дата должна иметь формат YYYY-MM-DD')
    if values.get('dateFrom') and values.get('dateTo') and values['dateFrom'] > values['dateTo']:
        fail(422, 'Начало периода позже окончания')
    if values.get('status') and values['status'] not in (ALLOCATION_STATUSES if assignment else PROJECT_STATUSES):
        fail(422, 'Неизвестный статус')
    return values


def latest_plans(c):
    result = {}
    for row in db.all_rows(c, 'plan'):
        plan = row['data']
        if plan['projectId'] not in result or plan['versionNumber'] > result[plan['projectId']]['versionNumber']:
            result[plan['projectId']] = plan
    return result


def visible_projects(c, a, plans):
    rows = []
    for row in db.all_rows(c, 'project'):
        p = row['data']
        if row['portfolio_id'] not in a.portfolios:continue
        if 'R-03' in a.roles and not a.roles & {'R-01', 'R-02'} and p['id'] not in a.projects and p['managerSubject'] != a.subject:continue
        if 'R-04' in a.roles and not a.roles & {'R-01', 'R-02', 'R-03'}:
            if not any(x['employeeId'] in a.employees for x in plans.get(p['id'], {}).get('allocations', [])):continue
        rows.append(p)
    return sorted(rows, key=lambda p: (p['priority'], p['code'], p['id']))


def project_rows(c, a, f):
    plans = latest_plans(c)
    rows = []
    for p in visible_projects(c, a, plans):
        v = plans.get(p['id'])
        if f.get('projectId') and p['id'] != f['projectId']:continue
        if f.get('portfolioId') and p['portfolioId'] != f['portfolioId']:continue
        if f.get('status') and p['status'] != f['status']:continue
        if f.get('q', '').casefold() not in (p['code']+' '+p['name']).casefold():continue
        if f.get('dateFrom') and (not v or v['horizonEnd'] < f['dateFrom']):continue
        if f.get('dateTo') and (not v or v['horizonStart'] > f['dateTo']):continue
        rows.append({'id':p['id'], 'code':p['code'], 'name':p['name'], 'status':p['status'],
                     'goal':p['goal'] if a.roles & {'R-01','R-02','R-03'} else '',
                     'portfolioId':p['portfolioId'], 'managerSubject':p['managerSubject'],
                     'versionNumber':v['versionNumber'] if v else None, 'planStatus':v['status'] if v else None,
                     'start':v['horizonStart'] if v else None, 'finish':v['horizonEnd'] if v else None})
    return rows


def employee_options(c, a, portfolio):
    identifiers = {r['data']['employeeId'] for r in db.all_rows(c, 'availability', portfolio)}
    if 'R-04' in a.roles and not a.roles & {'R-01','R-02','R-03'}:identifiers &= a.employees
    # Names may come from a local HR cache; availability IDs are the fallback.
    names = {r['id']:r['data'].get('displayName',r['id']) for r in db.all_rows(c, 'employee', portfolio)}
    local_names = {'employees/demo-1':'Демо-сотрудник 1', 'employees/demo-2':'Демо-сотрудник 2'}
    return [{'id':identifier, 'name':names.get(identifier, local_names.get(identifier, identifier))}
            for identifier in sorted(identifiers)]


def assignment_rows(c, a, f):
    plans = latest_plans(c)
    rows = []
    for p in visible_projects(c, a, plans):
        if f.get('projectId') and p['id'] != f['projectId']:continue
        if f.get('portfolioId') and p['portfolioId'] != f['portfolioId']:continue
        if f.get('q', '').casefold() not in (p['code']+' '+p['name']).casefold():continue
        v = plans.get(p['id'])
        if not v:continue
        names = {x['id']:x['name'] for x in employee_options(c,a,p['portfolioId'])}
        demands = {d['id']:d for d in v['demands']}
        items = {i['id']:i for i in v['items']}
        for allocation in v['allocations']:
            if 'R-04' in a.roles and not a.roles & {'R-01','R-02','R-03'} and allocation['employeeId'] not in a.employees:continue
            if f.get('employeeId') and allocation['employeeId'] != f['employeeId']:continue
            if f.get('status') and allocation['status'] != f['status']:continue
            if f.get('dateFrom') and allocation['date'] < f['dateFrom']:continue
            if f.get('dateTo') and allocation['date'] > f['dateTo']:continue
            demand = demands[allocation['demandId']]
            rows.append({**allocation, 'projectId':p['id'], 'projectCode':p['code'], 'projectName':p['name'],
                         'planVersionId':v['id'], 'versionNumber':v['versionNumber'], 'planStatus':v['status'],
                         'planItemId':demand['planItemId'], 'workName':items[demand['planItemId']]['name'],
                         'employeeName':names.get(allocation['employeeId'],allocation['employeeId'])})
    return sorted(rows, key=lambda r:(r['date'],r['projectCode'],r['workName'],r['employeeId'],r['id']))


def workbook_response(title, columns, rows, f, filename):
    if len(rows)>50000:fail(422, 'Слишком много строк: сузьте фильтры (максимум 50000)')
    book = Workbook()
    sheet = book.active
    sheet.title = title
    sheet.append([label for _,label in columns])
    for row in rows:
        cells=[]
        for key,_ in columns:
            value=row.get(key)
            if isinstance(value,str):value=ILLEGAL_CHARACTERS_RE.sub('',value)
            cell=Cell(sheet,value=value)
            # User text must never become an Excel formula or hyperlink.
            if isinstance(value,str):cell.data_type='s'
            if key in {'date','start','finish'} and value:
                cell.value=date.fromisoformat(value);cell.number_format='DD.MM.YYYY'
            if key=='hours':cell.number_format='0.00'
            cells.append(cell)
        sheet.append(cells)
    sheet.freeze_panes='A2'
    sheet.auto_filter.ref=sheet.dimensions
    for cell in sheet[1]:
        cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='235CCA')
    for index,(key,_) in enumerate(columns,1):
        sheet.column_dimensions[sheet.cell(1,index).column_letter].width=18 if key in {'date','start','finish','hours','status'} else 32
    meta=book.create_sheet('Фильтры')
    meta.append(['Параметр','Значение'])
    for key,value in [('Сформировано UTC',db.now()),('Строк',len(rows)),('Версии','Последняя версия каждого доступного проекта'),*sorted(f.items())]:
        if isinstance(value,str):value=ILLEGAL_CHARACTERS_RE.sub('',value)
        cells=[Cell(meta,value=str(key)),Cell(meta,value=value)]
        for cell in cells:
            if isinstance(cell.value,str):cell.data_type='s'
        meta.append(cells)
    meta.column_dimensions['A'].width=30;meta.column_dimensions['B'].width=65
    output=BytesIO();book.save(output)
    return Response(output.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition':f'attachment; filename="{filename}.xlsx"'})


def register(app, command_tx):
    from .app import allowed, get, changed, read_json

    @app.get('/api/v1/planning-options')
    def planning_options(req:Request):
        a=actor(req);allowed(a,['R-01','R-02','R-03','R-04'])
        with db.engine.begin() as c:
            portfolio=req.query_params.get('portfolioId','')
            get(c,'portfolio',portfolio,a)
            managers={r['data']['managerSubject']:r['data']['managerSubject'] for r in db.all_rows(c,'project',portfolio)}
            for r in db.all_rows(c,'manager',portfolio):
                managers[r['data']['subject']]=r['data']['displayName']
            if os.environ.get('APP_ENV')=='development':
                managers['demo-R-03']='Руководитель проекта · manager (учебная учётная запись)'
            if 'R-03' in a.roles:managers[a.subject]=managers.get(a.subject,'Моя учётная запись')
            if 'R-03' in a.roles and not a.roles & {'R-01','R-02'}:
                managers={a.subject:managers[a.subject]}
            return {'managers':[{'id':k,'name':v} for k,v in sorted(managers.items(),key=lambda x:x[1])],
                    'specialties':[{'id':k,'name':v} for k,v in [
                        ('developer','Разработчик'),('analyst','Аналитик'),('tester','Тестировщик'),
                        ('designer','Дизайнер'),('devops','DevOps-инженер'),('architect','Архитектор'),
                        ('project-manager','Руководитель проекта')]],
                    'grades':[{'id':k,'name':v} for k,v in [('intern','Стажёр'),('junior','Junior'),
                        ('middle','Middle'),('senior','Senior'),('lead','Ведущий специалист')]],
                    'source':'Локальный справочник Planning System; руководители — из проектов и кадрового кэша'}

    @app.patch('/api/v1/versions/{versionId}/timeline')
    async def change_timeline(req:Request):
        from .scheduling import reschedule
        a=actor(req);allowed(a,['R-02','R-03']);b=await read_json(req)
        if set(b)!={'itemId','start','finish'} or not all(isinstance(x,str) for x in b.values()):
            fail(422,'Нужны работа, начало и окончание')
        try:
            for key in ['start','finish']:
                if date.fromisoformat(b[key]).isoformat()!=b[key]:raise ValueError()
        except ValueError:fail(422,'Неверная дата')
        with command_tx(req,a,b) as (c,cached,complete):
            if cached:return cached
            v=get(c,'plan',req.path_params['versionId'],a);p=get(c,'project',v['projectId'],a)
            if v['status']!='DRAFT' or p['status']=='ARCHIVED':fail(409,'Сроки изменяются только в черновике действующего проекта')
            try:v,moved=reschedule(v,b['itemId'],b['start'],b['finish'])
            except (OverflowError,ValueError):fail(422,'Сроки вышли за допустимые границы')
            v['contentRevision']+=1;changed(c,'plan',v,req,a)
            return complete({'revision':v['revision'],'changedItemIds':moved},p['portfolioId'],p['id'])

    @app.get('/api/v1/project-timeline')
    def timeline(req:Request):
        a=actor(req);allowed(a,['R-01','R-02','R-03','R-04'])
        with db.engine.begin() as c:
            plans=latest_plans(c)
            pid=req.query_params.get('projectId','')
            if pid not in {p['id'] for p in visible_projects(c,a,plans)}:fail(404,'Проект не найден')
            v=plans.get(pid)
            return {k:v[k] for k in ['id','revision','status','items','dependencies','horizonStart','horizonEnd']} if v else {'items':[],'dependencies':[]}

    @app.get('/api/v1/project-cards')
    def cards(req:Request):
        a=actor(req);allowed(a,['R-01','R-02','R-03','R-04'])
        f=filters(req)
        with db.engine.begin() as c:
            rows=project_rows(c,a,f)
            return {'items':rows,'total':len(rows)}

    @app.get('/api/v1/assignments')
    def assignments(req:Request):
        a=actor(req);allowed(a,['R-01','R-02','R-03','R-04'])
        f=filters(req,True)
        with db.engine.begin() as c:
            rows=assignment_rows(c,a,f)
            return {'items':rows,'total':len(rows)}

    @app.get('/api/v1/resource-options')
    def resources(req:Request):
        a=actor(req);allowed(a,['R-01','R-02','R-03','R-04'])
        with db.engine.begin() as c:
            p=get(c,'project',req.query_params.get('projectId',''),a)
            v=latest_plans(c).get(p['id'])
            if p['id'] not in {x['id'] for x in visible_projects(c,a,latest_plans(c))}:fail(404,'Проект не найден')
            employees=employee_options(c,a,p['portfolioId'])
            ids={x['id'] for x in employees}
            items={x['id']:x for x in v['items']} if v else {}
            return {'employees':employees,'versionId':v['id'] if v else None,'revision':v['revision'] if v else None,
                    'planStatus':v['status'] if v else None,
                    'demands':[dict(d,workName=items[d['planItemId']]['name']) for d in v['demands']] if v else [],
                    'availability':[{'data':r['data'],'revision':r['revision']} for r in db.all_rows(c,'availability',p['portfolioId']) if r['data']['employeeId'] in ids]}

    @app.post('/api/v1/versions/{versionId}/assignments',status_code=201)
    @app.patch('/api/v1/versions/{versionId}/assignments/{assignmentId}')
    async def add_assignment(req:Request):
        a=actor(req);allowed(a,['R-02','R-04']);b=await read_json(req)
        editing='assignmentId' in req.path_params
        if not {'demandId','employeeId','date','hours'}<=set(b) or set(b)-{'demandId','employeeId','date','hours','status'}:
            fail(422,'Нужны работа, сотрудник, дата и часы')
        if b.get('status','PROPOSED') not in {'PROPOSED','CANCELLED'} or (not editing and 'status' in b):
            fail(422,'Подтверждение выполняет ресурсный руководитель отдельной операцией')
        if isinstance(b['hours'],bool) or not isinstance(b['hours'],(int,float)) or not math.isfinite(b['hours']) or b['hours']<=0 or round(b['hours'],2)!=b['hours']:
            fail(422,'Часы должны быть положительными с точностью до сотых')
        if not all(isinstance(b[k],str) for k in ['demandId','employeeId','date']):fail(422,'Неверный формат назначения')
        try:
            if date.fromisoformat(b['date']).isoformat()!=b['date']:raise ValueError()
        except ValueError:fail(422,'Неверная дата назначения')
        with command_tx(req,a,b) as (c,cached,complete):
            if cached:return cached
            v=get(c,'plan',req.path_params['versionId'],a);p=get(c,'project',v['projectId'],a)
            if p['status']=='ARCHIVED' or v['status']!='DRAFT':fail(409,'Назначения добавляются только в черновик действующего проекта')
            if b['employeeId'] not in {x['id'] for x in employee_options(c,a,p['portfolioId'])}:fail(403,'Сотрудник недоступен')
            if 'R-04' in a.roles and 'R-02' not in a.roles and b['employeeId'] not in a.employees:fail(403,'Чужой сотрудник')
            demand=next((d for d in v['demands'] if d['id']==b['demandId']),None)
            if not demand:fail(422,'Потребность не принадлежит этой версии')
            previous=next((x for x in v['allocations'] if x['id']==req.path_params.get('assignmentId')),None)
            if editing and not previous:fail(404,'Назначение не найдено')
            if editing and 'R-04' in a.roles and 'R-02' not in a.roles and previous['employeeId'] not in a.employees:fail(403,'Чужое назначение')
            if b.get('status')!='CANCELLED' and any(x['id']!=req.path_params.get('assignmentId') and x['demandId']==b['demandId'] and x['employeeId']==b['employeeId'] and x['date']==b['date'] and x['status']!='CANCELLED' for x in v['allocations']):
                fail(409,'На эту работу и дату сотрудник уже назначен')
            allocation={'id':previous['id'] if editing else db.uid(),**b,'status':b.get('status','PROPOSED')}
            v=copy.deepcopy(v)
            if editing:v['allocations']=[allocation if x['id']==allocation['id'] else x for x in v['allocations']]
            else:v['allocations'].append(allocation)
            validate_draft(v)
            v['contentRevision']+=1;changed(c,'plan',v,req,a)
            return complete({'assignment':allocation,'revision':v['revision']},p['portfolioId'],p['id'],status=200 if editing else 201)

    @app.get('/api/v1/downloads/projects.xlsx')
    def export_projects(req:Request):
        a=actor(req);allowed(a,['R-01','R-02','R-03'])
        f=filters(req)
        with db.engine.begin() as c:
            rows=project_rows(c,a,f)
            return workbook_response('Проекты',[('code','Код'),('name','Проект'),('status','Статус проекта'),('goal','Цель'),
                ('managerSubject','Руководитель'),('versionNumber','Версия плана'),('planStatus','Статус плана'),('start','Начало'),('finish','Окончание')],rows,f,'projects')

    @app.get('/api/v1/downloads/assignments.xlsx')
    def export_assignments(req:Request):
        a=actor(req);allowed(a,['R-01','R-02','R-03','R-04'])
        f=filters(req,True)
        with db.engine.begin() as c:
            rows=assignment_rows(c,a,f)
            return workbook_response('Назначения',[('projectCode','Код проекта'),('projectName','Проект'),('versionNumber','Версия'),
                ('planStatus','Статус плана'),('workName','Работа'),('employeeName','Сотрудник'),('employeeId','ID сотрудника'),
                ('date','Дата'),('hours','Часы'),('status','Статус назначения')],rows,f,'assignments')
